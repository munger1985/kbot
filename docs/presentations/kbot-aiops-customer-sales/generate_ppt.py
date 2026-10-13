from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "KBot_AIOps_客户与销售介绍_2026-10.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = "10243A"
NAVY = "17324D"
BLUE = "2C6BED"
TEAL = "00A69A"
ORANGE = "F2A93B"
RED = "D9534F"
WHITE = "FFFFFF"
PAPER = "F4F7FA"
PAPER_2 = "EAF0F5"
MUTED = "627386"
MID = "9AA8B5"
LINE = "D8E1E8"
LIGHT_BLUE = "E9F0FF"
LIGHT_TEAL = "E5F7F4"
LIGHT_ORANGE = "FFF3DC"
LIGHT_RED = "FDEBEC"

FONT = "Microsoft YaHei"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_bg(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin: float = 0.04,
    font: str = FONT,
    line_spacing: float | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    for current in frame.paragraphs:
        current.alignment = align
        if line_spacing is not None:
            current.line_spacing = line_spacing
        for run in current.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, runs, x, y, w, h, *, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for item in runs:
        run = paragraph.add_run()
        run.text = item[0]
        run.font.name = FONT
        run.font.size = Pt(item[1])
        run.font.bold = item[2]
        run.font.color.rgb = rgb(item[3])
    return box


def add_shape(
    slide,
    shape_type,
    x,
    y,
    w,
    h,
    *,
    fill=WHITE,
    line=LINE,
    line_width=1.0,
    radius=False,
):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(line_width)
    return shape


def add_round_rect(slide, x, y, w, h, *, fill=WHITE, line=LINE, line_width=1.0):
    return add_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x,
        y,
        w,
        h,
        fill=fill,
        line=line,
        line_width=line_width,
    )


def add_line(slide, x1, y1, x2, y2, *, color=LINE, width=1.5, dash=None):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    if dash is not None:
        line.line.dash_style = dash
    return line


def add_pill(slide, text, x, y, w, *, fill=LIGHT_BLUE, color=BLUE, line=None, size=11):
    add_round_rect(slide, x, y, w, 0.34, fill=fill, line=line or fill, line_width=0.5)
    return add_text(slide, text, x, y, w, 0.34, size=size, color=color, bold=True, align=PP_ALIGN.CENTER)


def add_kbot_mark(slide, *, dark=False):
    color = WHITE if dark else INK
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.55, 0.40, 0.08, 0.35, fill=TEAL, line=TEAL, line_width=0)
    add_text(slide, "KBot", 0.72, 0.35, 0.9, 0.45, size=17, color=color, bold=True)


def add_footer(slide, page: int, *, dark=False, source: str | None = None):
    color = "B9C8D6" if dark else MUTED
    add_text(slide, "KBot AIOps · 客户与销售介绍", 0.62, 7.08, 4.2, 0.22, size=8.5, color=color)
    if source:
        add_text(slide, source, 4.45, 7.08, 7.65, 0.22, size=7.5, color=color, align=PP_ALIGN.RIGHT)
    add_text(slide, f"{page:02d}", 12.18, 7.03, 0.5, 0.25, size=9, color=color, bold=True, align=PP_ALIGN.RIGHT)


def add_title(slide, kicker: str, title: str, subtitle: str | None = None, *, dark=False):
    fg = WHITE if dark else INK
    muted = "B9C8D6" if dark else MUTED
    add_kbot_mark(slide, dark=dark)
    add_text(slide, kicker.upper(), 0.62, 0.95, 3.5, 0.28, size=10, color=TEAL, bold=True)
    add_text(slide, title, 0.62, 1.18, 11.8, 0.64, size=30, color=fg, bold=True, valign=MSO_ANCHOR.TOP)
    if subtitle:
        add_text(slide, subtitle, 0.66, 1.82, 11.4, 0.34, size=13.5, color=muted, valign=MSO_ANCHOR.TOP)


def add_icon_disc(slide, label: str, x, y, *, fill=BLUE, size=0.54, text_color=WHITE, font_size=15):
    add_shape(slide, MSO_SHAPE.OVAL, x, y, size, size, fill=fill, line=fill, line_width=0)
    add_text(slide, label, x, y, size, size, size=font_size, color=text_color, bold=True, align=PP_ALIGN.CENTER)


def add_card(slide, x, y, w, h, *, title, body, number=None, accent=BLUE, fill=WHITE):
    add_round_rect(slide, x, y, w, h, fill=fill, line=LINE, line_width=0.8)
    if number is not None:
        add_text(slide, number, x + 0.22, y + 0.16, 0.55, 0.32, size=10, color=accent, bold=True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, x + 0.22, y + 0.55, 0.06, 0.50, fill=accent, line=accent, line_width=0)
    add_text(slide, title, x + 0.42, y + 0.50, w - 0.65, 0.48, size=17, color=INK, bold=True)
    add_text(slide, body, x + 0.42, y + 1.03, w - 0.68, h - 1.20, size=12.5, color=MUTED, valign=MSO_ANCHOR.TOP, line_spacing=1.1)


def add_stat(slide, value, label, x, y, w, *, accent=TEAL, dark=False):
    fill = NAVY if dark else WHITE
    line = "36526D" if dark else LINE
    fg = WHITE if dark else INK
    muted = "B9C8D6" if dark else MUTED
    add_round_rect(slide, x, y, w, 0.95, fill=fill, line=line, line_width=0.8)
    add_text(slide, value, x + 0.18, y + 0.08, 1.05, 0.52, size=25, color=accent, bold=True)
    add_text(slide, label, x + 1.18, y + 0.10, w - 1.35, 0.50, size=11.5, color=fg, bold=True)
    add_text(slide, "当前代码目录", x + 1.18, y + 0.56, w - 1.35, 0.20, size=8.5, color=muted)


def slide_01(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, INK)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.16, 7.5, fill=TEAL, line=TEAL, line_width=0)
    add_kbot_mark(slide, dark=True)
    add_text(slide, "KBot AIOps", 0.72, 1.38, 5.9, 0.78, size=38, color=WHITE, bold=True)
    add_text(slide, "可信、可控的数据库智能运维", 0.72, 2.15, 5.9, 0.72, size=28, color=WHITE, bold=True, valign=MSO_ANCHOR.TOP)
    add_text(
        slide,
        "把监控、数据库、运维知识与人的判断\n组织成一条可验证的闭环。",
        0.76,
        3.08,
        5.15,
        1.05,
        size=17,
        color="C5D2DE",
        valign=MSO_ANCHOR.TOP,
        line_spacing=1.15,
    )
    add_pill(slide, "客户与销售介绍", 0.76, 5.48, 1.55, fill="24445F", color=WHITE, line="34546E", size=10)
    add_text(slide, "2026 · 基于当前 KBot4 能力", 0.78, 5.92, 3.2, 0.24, size=10, color="9FB1C2")

    cx, cy = 9.55, 3.52
    nodes = [
        (9.45, 0.98, "发现", BLUE),
        (11.10, 1.80, "观测", TEAL),
        (11.35, 3.70, "诊断", BLUE),
        (10.06, 5.12, "验证", TEAL),
        (8.10, 4.88, "报告", ORANGE),
        (7.55, 2.95, "处置", ORANGE),
    ]
    for nx, ny, _, _ in nodes:
        add_line(slide, cx, cy, nx + 0.45, ny + 0.45, color="35516A", width=1.4)
    add_shape(slide, MSO_SHAPE.OVAL, 8.65, 2.62, 1.8, 1.8, fill="1E4662", line=TEAL, line_width=2.0)
    add_text(slide, "AIOps", 8.65, 2.84, 1.8, 0.54, size=21, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "证据闭环", 8.65, 3.40, 1.8, 0.30, size=10.5, color="BFD0DE", align=PP_ALIGN.CENTER)
    for nx, ny, label, accent in nodes:
        add_shape(slide, MSO_SHAPE.OVAL, nx, ny, 0.9, 0.9, fill=NAVY, line=accent, line_width=1.8)
        add_text(slide, label, nx, ny, 0.9, 0.9, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 1, dark=True)


def slide_02(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PAPER)
    add_title(slide, "01 · WHAT", "什么是 AIOps？", "不是“让 AI 接管运维”，而是让 AI 在真实数据和治理边界内参与决策。")
    add_rich_text(
        slide,
        [("AIOps = ", 27, True, INK), ("AI 决策能力", 27, True, BLUE), (" × 运维数据 × 治理闭环", 27, True, INK)],
        0.68,
        2.32,
        11.8,
        0.72,
    )
    cards = [
        ("AI 决策能力", "理解问题、选择诊断路径、形成可解释假设与建议。", "AI", BLUE),
        ("运维数据", "连接监控、数据库、日志、知识与人工补充证据。", "DATA", TEAL),
        ("治理闭环", "权限、策略、审批、执行复核、效果验证与审计。", "CTRL", ORANGE),
    ]
    x_positions = [0.68, 4.50, 8.32]
    for x, (title, body, label, accent) in zip(x_positions, cards):
        add_round_rect(slide, x, 3.35, 3.40, 2.12, fill=WHITE, line=LINE, line_width=0.8)
        add_icon_disc(slide, label, x + 0.25, 3.62, fill=accent, size=0.66, font_size=9 if label != "AI" else 15)
        add_text(slide, title, x + 1.05, 3.55, 2.05, 0.55, size=17, color=INK, bold=True)
        add_text(slide, body, x + 0.27, 4.34, 2.85, 0.82, size=12.5, color=MUTED, valign=MSO_ANCHOR.TOP, line_spacing=1.1)
    add_round_rect(slide, 0.68, 5.78, 11.04, 0.70, fill=INK, line=INK, line_width=0)
    add_text(slide, "核心价值：从“告警发生了”走到“为什么、怎么办、做完是否有效”。", 0.95, 5.78, 10.5, 0.70, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 2)


def slide_03(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PAPER)
    add_title(slide, "02 · WHY", "传统运维的难点，不是缺工具，而是缺一条共同的判断链", "监控看见现象，专家解释原因，流程控制行动——三者往往彼此割裂。")
    add_card(slide, 0.68, 2.40, 2.75, 2.55, title="信息分散", body="指标、日志、数据库状态和手册分布在不同系统，排障靠人工拼接。", number="01", accent=BLUE)
    add_card(slide, 3.62, 2.40, 2.75, 2.55, title="告警没有答案", body="告警说明“异常发生”，却不说明影响、根因、证据缺口和下一步。", number="02", accent=TEAL)
    add_card(slide, 6.56, 2.40, 2.75, 2.55, title="经验难复制", body="关键判断依赖少数专家，过程没有标准结构，也难复用到下一次。", number="03", accent=ORANGE)
    add_card(slide, 9.50, 2.40, 2.75, 2.55, title="处置难审计", body="建议、批准、执行和验证缺少统一记录，命令成功也不代表问题解决。", number="04", accent=RED)
    add_round_rect(slide, 1.18, 5.35, 10.05, 0.88, fill=WHITE, line=LINE, line_width=0.8)
    add_text(slide, "KBot AIOps 把一次排障变成可复盘、可复用、可治理的组织资产。", 1.48, 5.35, 9.45, 0.88, size=18.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 3)


def slide_04(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, INK)
    add_title(slide, "03 · LOOP", "从发现问题，到证明问题已经改善", "同一条闭环承接人工对话、监控告警和定时巡检。", dark=True)
    steps = [
        ("01", "发现", "对话 / 告警 / 巡检", BLUE),
        ("02", "观测", "监控与数据库证据", TEAL),
        ("03", "诊断", "候选根因与反证", BLUE),
        ("04", "评估", "证据等级与缺口", ORANGE),
        ("05", "建议", "缓解与长期方案", TEAL),
        ("06", "审批", "逐条人工把关", ORANGE),
        ("07", "执行", "受控动作目录", RED),
        ("08", "验证", "同口径 Before / After", TEAL),
        ("09", "报告", "复盘、巡检与对比", BLUE),
    ]
    x0, y0, w, gap = 0.68, 2.55, 1.24, 0.10
    for i, (num, title, desc, accent) in enumerate(steps):
        x = x0 + i * (w + gap)
        add_round_rect(slide, x, y0, w, 2.18, fill=NAVY, line="36526D", line_width=0.7)
        add_text(slide, num, x + 0.12, y0 + 0.12, 0.48, 0.28, size=9.5, color=accent, bold=True)
        add_text(slide, title, x + 0.12, y0 + 0.48, w - 0.24, 0.48, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x + 0.12, y0 + 1.05, w - 0.24, 0.82, size=10.2, color="B9C8D6", align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
        if i < len(steps) - 1:
            add_text(slide, "›", x + w - 0.01, y0 + 0.78, gap + 0.12, 0.45, size=22, color="7790A7", bold=True, align=PP_ALIGN.CENTER)
    add_round_rect(slide, 1.25, 5.35, 10.83, 0.72, fill="1E4662", line="2C6387", line_width=0.7)
    add_text(slide, "每一步都保存事实、责任人与版本；无法判断时，系统如实说明证据不足。", 1.55, 5.35, 10.25, 0.72, size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 4, dark=True)


def slide_05(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PAPER)
    add_title(slide, "04 · CAPABILITIES", "一套平台，连接数据、形成判断、控制行动", "能力不是孤立菜单，而是围绕运维任务（Run）协同工作。")
    layers = [
        ("连接层", "运维目标 · 监控源 · 数据库 · 日志 · 文件 · 运维知识", 2.42, LIGHT_BLUE, BLUE),
        ("智能层", "私有 DBA Agent · 诊断工具 · Playbook · 根因评估 · 人工补证", 3.36, LIGHT_TEAL, TEAL),
        ("治理层", "策略 · 权限 · 审批 · 受控执行 · 处理后验证 · 版本化报告", 4.30, LIGHT_ORANGE, ORANGE),
    ]
    for label, body, y, fill, accent in layers:
        add_round_rect(slide, 0.70, y, 8.15, 0.72, fill=fill, line=fill, line_width=0)
        add_round_rect(slide, 0.88, y + 0.13, 1.12, 0.44, fill=accent, line=accent, line_width=0)
        add_text(slide, label, 0.88, y + 0.13, 1.12, 0.44, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, 2.22, y + 0.08, 6.27, 0.55, size=13, color=INK, bold=True)
    add_stat(slide, "154", "版本化只读诊断工具\nOracle · MySQL · PostgreSQL", 9.15, 2.42, 3.12, accent=BLUE)
    add_stat(slide, "34", "审批后可执行动作模板\nOracle · MySQL", 9.15, 3.58, 3.12, accent=ORANGE)
    add_stat(slide, "3", "主要业务入口\n对话 · 告警 · 巡检", 9.15, 4.74, 3.12, accent=TEAL)
    add_text(slide, "边界说明：动作是否可执行，仍取决于授权、策略、部署开关、凭据与人工审批。", 0.75, 5.56, 8.2, 0.55, size=11, color=MUTED)
    add_footer(slide, 5, source="数据源：当前诊断与动作目录（2026-10-12）")


def slide_06(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PAPER)
    add_title(slide, "05 · ENTRY POINTS", "三种入口，进入同一条诊断与报告链路", "无论问题从哪里开始，证据、策略、Agent 版本和审计方式保持一致。")
    entries = [
        ("对话诊断", "即时排障与多轮追问", "适合值班人员、DBA 与应用负责人主动提问。", "CHAT", BLUE),
        ("告警诊断", "监控事件自动创建 Run", "验签、映射、去重后启动诊断；证据不足则输出受限结论。", "ALERT", ORANGE),
        ("定时巡检", "日报、周报与 Cron 计划", "按计划冻结 Agent 与 Target 范围，持续沉淀风险和趋势。", "PLAN", TEAL),
    ]
    xs = [0.68, 4.40, 8.12]
    for x, (title, subtitle, body, label, accent) in zip(xs, entries):
        add_round_rect(slide, x, 2.36, 3.32, 2.56, fill=WHITE, line=LINE, line_width=0.8)
        add_icon_disc(slide, label, x + 0.24, 2.66, fill=accent, size=0.65, font_size=8.5)
        add_text(slide, title, x + 1.06, 2.58, 1.95, 0.42, size=18, color=INK, bold=True)
        add_text(slide, subtitle, x + 0.26, 3.40, 2.78, 0.36, size=12, color=accent, bold=True)
        add_text(slide, body, x + 0.26, 3.88, 2.76, 0.78, size=11.7, color=MUTED, valign=MSO_ANCHOR.TOP)
        add_line(slide, x + 1.66, 4.92, 6.68, 5.42, color=MID, width=1.2, dash=MSO_LINE_DASH_STYLE.DASH)
    add_round_rect(slide, 3.26, 5.15, 6.82, 0.82, fill=INK, line=INK, line_width=0)
    add_text(slide, "统一运维任务：证据 → 判断 → 建议 → 验证 → 报告", 3.55, 5.15, 6.24, 0.82, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 6)


def slide_07(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PAPER)
    add_title(slide, "06 · EVIDENCE", "KBot 的特色：AI 结论必须带着证据等级", "知识可以解释机制，模型可以提出假设，但现场事实必须来自可追溯证据。")
    evidence = [
        ("系统验证", "监控或数据库直接观测", BLUE),
        ("用户提供", "SQL 结果、文本、CSV、截图", TEAL),
        ("知识引用", "SOP、手册与历史经验", ORANGE),
        ("模型推断", "摘要、假设与因果推理", MUTED),
    ]
    for i, (title, body, accent) in enumerate(evidence):
        y = 2.36 + i * 0.88
        add_round_rect(slide, 0.68, y, 3.22, 0.68, fill=WHITE, line=LINE, line_width=0.7)
        add_icon_disc(slide, str(i + 1), 0.88, y + 0.10, fill=accent, size=0.48, font_size=11)
        add_text(slide, title, 1.52, y + 0.06, 0.96, 0.29, size=13, color=INK, bold=True)
        add_text(slide, body, 1.52, y + 0.35, 2.05, 0.23, size=9.8, color=MUTED)
    add_round_rect(slide, 4.45, 2.66, 3.05, 3.02, fill=INK, line=INK, line_width=0)
    add_text(slide, "候选根因", 4.82, 2.92, 2.32, 0.42, size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, 4.92, 3.60, 7.02, 3.60, color="496981", width=1.2)
    add_text(slide, "支持证据", 4.80, 3.78, 1.08, 0.33, size=12, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "反证", 6.03, 3.78, 1.08, 0.33, size=12, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "数据缺口", 5.42, 4.38, 1.08, 0.33, size=12, color="C5D2DE", bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "证据评估", 4.82, 5.02, 2.32, 0.36, size=13, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    levels = [
        ("已确认", "直接证据充分", TEAL, LIGHT_TEAL),
        ("很可能", "多项证据一致", BLUE, LIGHT_BLUE),
        ("可能", "仅解释部分现象", ORANGE, LIGHT_ORANGE),
        ("无法判断", "证据不足或冲突", RED, LIGHT_RED),
    ]
    for i, (title, body, accent, fill) in enumerate(levels):
        y = 2.36 + i * 0.88
        add_round_rect(slide, 8.05, y, 4.18, 0.68, fill=fill, line=fill, line_width=0)
        add_text(slide, title, 8.28, y + 0.07, 1.15, 0.29, size=13, color=accent, bold=True)
        add_text(slide, body, 9.52, y + 0.08, 2.38, 0.28, size=11, color=INK, bold=True)
    add_footer(slide, 7)


def slide_08(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PAPER)
    add_title(slide, "07 · COLLABORATION", "数据不完整时，人和 AI 仍能继续工作", "KBot 不用“猜”补齐现场，而是发起最小补证请求，并把人的输入纳入同一证据链。")
    flow = [
        ("1", "发现缺口", "缺少关键会话、SQL 或现场截图", BLUE),
        ("2", "提出请求", "生成范围受控的只读 SQL 或材料清单", ORANGE),
        ("3", "人工补证", "用户在有权限的环境执行并回贴", TEAL),
        ("4", "继续分析", "原 Run 恢复，重新评估支持与反证", BLUE),
        ("5", "受限结论", "仍不足时明确说明，不制造确定性", RED),
    ]
    for i, (num, title, body, accent) in enumerate(flow):
        x = 0.70 + i * 2.34
        add_icon_disc(slide, num, x, 2.55, fill=accent, size=0.52, font_size=12)
        add_text(slide, title, x + 0.66, 2.47, 1.48, 0.35, size=14, color=INK, bold=True)
        add_text(slide, body, x, 3.18, 1.98, 0.86, size=10.5, color=MUTED, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
        if i < len(flow) - 1:
            add_text(slide, "→", x + 1.96, 2.57, 0.34, 0.34, size=18, color=MID, bold=True, align=PP_ALIGN.CENTER)
    add_round_rect(slide, 0.70, 4.48, 11.54, 1.22, fill=INK, line=INK, line_width=0)
    add_text(slide, "同一平台可配置多个私有 DBA Agent", 1.03, 4.67, 3.30, 0.34, size=16, color=WHITE, bold=True)
    agents = [("核心生产库", "严格策略", ORANGE), ("测试环境", "快速诊断", BLUE), ("值班排障", "告警联动", TEAL), ("专项巡检", "固定范围", "91A6B7")]
    for i, (name, desc, accent) in enumerate(agents):
        x = 4.40 + i * 1.82
        add_round_rect(slide, x, 4.68, 1.55, 0.72, fill=NAVY, line="35526A", line_width=0.7)
        add_text(slide, name, x, 4.71, 1.55, 0.30, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x, 5.03, 1.55, 0.23, size=8.5, color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 8)


def slide_09(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, INK)
    add_title(slide, "08 · CONTROL", "让 AI 参与处置，但不让 AI 绕过责任边界", "任何可执行动作都要通过多道治理闸门。", dark=True)
    gates = [
        ("授权", "Agent 与 Target\n显式绑定", BLUE),
        ("策略", "环境、风险与\n动作范围", TEAL),
        ("审批", "每条命令\n独立批准", ORANGE),
        ("授权令牌", "一次性、短时、\n绑定内容摘要", BLUE),
        ("执行复核", "模板、参数与\n前置条件校验", RED),
        ("效果验证", "同口径采集\n主指标与护栏", TEAL),
    ]
    for i, (title, body, accent) in enumerate(gates):
        x = 0.65 + i * 2.08
        add_shape(slide, MSO_SHAPE.HEXAGON, x, 2.55, 1.78, 1.78, fill=NAVY, line=accent, line_width=1.6)
        add_text(slide, title, x + 0.16, 2.83, 1.46, 0.38, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.16, 3.22, 1.46, 0.68, size=9.8, color="B9C8D6", align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
        if i < len(gates) - 1:
            add_text(slide, "›", x + 1.75, 3.12, 0.34, 0.45, size=21, color="6C849A", bold=True, align=PP_ALIGN.CENTER)
    add_round_rect(slide, 0.95, 4.95, 11.18, 0.98, fill="1E4662", line="2C6387", line_width=0.7)
    add_text(slide, "已登记动作示例", 1.22, 5.14, 1.58, 0.30, size=12, color=TEAL, bold=True)
    add_text(slide, "会话控制 · 索引与统计维护 · Scheduler · 存储增长 · 动态参数 · 受限权限维护", 2.86, 5.06, 8.78, 0.46, size=13.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "PostgreSQL 当前以只读诊断为主；破坏性动作保持人工执行或不支持。", 1.20, 6.22, 10.9, 0.30, size=10.5, color="AFC0CE", align=PP_ALIGN.CENTER)
    add_footer(slide, 9, dark=True, source="实际可执行范围以 Action Catalog、授权与部署开关为准")


def slide_10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PAPER)
    add_title(slide, "09 · OUTPUT", "每次处理都会留下可交付、可复盘的结果", "报告不是 AI 临时作文，而是从已保存的事实、证据、审批与验证结果生成。")
    reports = [
        ("故障报告", "时间线 · 影响 · 根因 · 处置", RED),
        ("性能报告", "瓶颈 · 关键指标 · 优化建议", BLUE),
        ("巡检报告", "健康状态 · 风险 · 待办", TEAL),
        ("工作负载", "AWR / ASH / Workload / Diff", ORANGE),
        ("前后对比", "主指标 · 护栏 · 最终结论", BLUE),
    ]
    for i, (title, body, accent) in enumerate(reports):
        x = 0.68 + i * 2.37
        add_round_rect(slide, x, 2.42, 2.08, 1.78, fill=WHITE, line=LINE, line_width=0.8)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.42, 2.08, 0.10, fill=accent, line=accent, line_width=0)
        add_text(slide, title, x + 0.18, 2.72, 1.72, 0.38, size=15.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.18, 3.22, 1.72, 0.54, size=10.3, color=MUTED, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
    add_text(slide, "对客户的业务价值", 0.72, 4.67, 2.5, 0.38, size=16, color=INK, bold=True)
    values = [
        ("更快进入有效分析", "少在工具之间搬运数据，把时间用于判断。", BLUE),
        ("把专家经验变成流程", "诊断工具、Playbook 和证据规则可以持续复用。", TEAL),
        ("让每次行动可解释", "谁判断、谁批准、执行了什么、是否改善，都有记录。", ORANGE),
    ]
    for i, (title, body, accent) in enumerate(values):
        x = 0.72 + i * 4.0
        add_icon_disc(slide, str(i + 1), x, 5.22, fill=accent, size=0.48, font_size=11)
        add_text(slide, title, x + 0.64, 5.12, 2.85, 0.36, size=14.2, color=INK, bold=True)
        add_text(slide, body, x + 0.64, 5.54, 3.00, 0.68, size=10.7, color=MUTED, valign=MSO_ANCHOR.TOP)
    add_footer(slide, 10)


def slide_11(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PAPER)
    add_title(slide, "10 · DEPLOYMENT", "单机一体化起步，按规模平滑拆分", "下图是面向 PoC 与中小规模生产的推荐方案；模型走 API 或独立推理节点。")

    add_round_rect(slide, 0.66, 2.34, 1.62, 1.10, fill=WHITE, line=LINE, line_width=0.8)
    add_text(slide, "客户用户", 0.84, 2.50, 1.25, 0.30, size=14, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "浏览器 / 销售演示", 0.84, 2.86, 1.25, 0.24, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "→", 2.32, 2.65, 0.45, 0.38, size=22, color=BLUE, bold=True, align=PP_ALIGN.CENTER)

    add_round_rect(slide, 2.78, 2.18, 6.55, 3.48, fill=WHITE, line=BLUE, line_width=1.4)
    add_pill(slide, "一台 Linux x86_64 服务器", 3.05, 2.34, 2.33, fill=INK, color=WHITE, line=INK, size=10)
    components = [
        ("入口层", "Nginx / TLS\nORDS + Oracle APEX", 3.05, 2.92, LIGHT_BLUE, BLUE),
        ("数据层", "Oracle Database 26ai\nKBot Schema / 审计事实", 5.10, 2.92, LIGHT_ORANGE, ORANGE),
        ("应用层", "Python 3.10+ / FastAPI\nKBot API · Worker · Scheduler", 7.15, 2.92, LIGHT_TEAL, TEAL),
        ("可选观测", "Prometheus · Loki · Alloy\nExporter / Collector", 3.05, 4.20, "EDF2F6", MUTED),
        ("模型接入", "本地 OCR / Embedding\n近端 DeepSeek 或云 API", 5.10, 4.20, LIGHT_BLUE, BLUE),
        ("安全运维", "Secret · 最小权限\n日志 · 备份 · 进程编排", 7.15, 4.20, "EDF2F6", MUTED),
    ]
    for title, body, x, y, fill, accent in components:
        add_round_rect(slide, x, y, 1.78, 0.96, fill=fill, line=fill, line_width=0)
        add_text(slide, title, x + 0.10, y + 0.08, 1.58, 0.25, size=11, color=accent, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.10, y + 0.36, 1.58, 0.47, size=8.7, color=INK, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)

    add_text(slide, "→", 9.42, 2.72, 0.45, 0.38, size=22, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_round_rect(slide, 9.92, 2.14, 2.66, 1.54, fill=WHITE, line=LINE, line_width=0.8)
    add_text(slide, "运维目标", 10.18, 2.34, 2.12, 0.30, size=14, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Oracle 19c / 26ai\nMySQL 8.4 · PostgreSQL 16", 10.18, 2.78, 2.12, 0.58, size=10.5, color=MUTED, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
    add_round_rect(slide, 9.92, 4.12, 2.66, 1.54, fill=WHITE, line=LINE, line_width=0.8)
    add_text(slide, "外部服务", 10.18, 4.31, 2.12, 0.30, size=14, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Prometheus / Zabbix / OEM\nOCI / 模型 API / 近端推理", 10.18, 4.73, 2.12, 0.58, size=10.2, color=MUTED, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)

    add_round_rect(slide, 0.66, 5.96, 11.92, 0.72, fill=INK, line=INK, line_width=0)
    specs = [("16 vCPU", "计算"), ("64 GB RAM", "内存"), ("1 TB NVMe", "存储起步"), ("GPU 可选", "大模型建议独立节点")]
    for i, (value, label) in enumerate(specs):
        x = 0.92 + i * 2.78
        add_text(slide, value, x, 6.03, 1.52, 0.28, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 1.50, 6.05, 1.10, 0.24, size=8.5, color="AFC0CE", align=PP_ALIGN.LEFT)
    add_text(slide, "推荐起步配置；高可用、大规模目标或本地大模型应拆分数据库、应用、观测与推理节点。", 0.80, 6.74, 11.7, 0.22, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_footer(slide, 11, source="硬件规格为项目建议，不是代码声明的最低要求")


def slide_12(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PAPER)
    add_title(slide, "11 · MODELS", "KBot 支持的模型与接入方式", "模型目录默认不预置业务模型；实际启用项由客户部署后的 Model Catalog 决定。")
    columns = [0.68, 2.28, 7.55, 10.00]
    widths = [1.48, 5.15, 2.32, 2.28]
    headers = ["能力类别", "支持的模型 / Provider", "接入形态", "开放性说明"]
    for x, w, title in zip(columns, widths, headers):
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 2.30, w, 0.50, fill=INK, line=INK, line_width=0)
        add_text(slide, title, x + 0.05, 2.30, w - 0.10, 0.50, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    rows = [
        ("LLM", "DeepSeek 近端 / DeepSeek API / Qwen API / ChatGPT / OCI", "近端 API\n云 API", "开放权重近端\n+ 闭源服务", TEAL),
        ("文本向量", "BGE Local / Qwen3 Embedding Local / Qwen API / OpenAI / OCI", "本地加载\n云 API", "开源或开放权重\n+ 闭源服务", BLUE),
        ("视觉向量", "ColQwen2 Local", "本地加载", "开放权重", ORANGE),
        ("VLM", "Qwen-VL API / GPT-4V 兼容接口", "云 API", "API 服务；\nGPT 为闭源", BLUE),
        ("OCR", "RapidOCR / EasyOCR / Tesseract / DeepSeek OCR API", "本地引擎\n云 API", "开源本地\n+ API 服务", TEAL),
    ]
    y = 2.80
    row_h = 0.72
    for i, (category, models, access, openness, accent) in enumerate(rows):
        fill = WHITE if i % 2 == 0 else "EDF2F6"
        for x, w in zip(columns, widths):
            add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, row_h, fill=fill, line=LINE, line_width=0.55)
        add_pill(slide, category, columns[0] + 0.18, y + 0.19, widths[0] - 0.36, fill=accent, color=WHITE, line=accent, size=10)
        add_text(slide, models, columns[1] + 0.16, y + 0.06, widths[1] - 0.30, row_h - 0.12, size=10.3, color=INK, bold=True)
        add_text(slide, access, columns[2] + 0.10, y + 0.05, widths[2] - 0.20, row_h - 0.10, size=9.8, color=MUTED, align=PP_ALIGN.CENTER)
        add_text(slide, openness, columns[3] + 0.10, y + 0.05, widths[3] - 0.20, row_h - 0.10, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
        y += row_h

    add_round_rect(slide, 0.68, 6.62, 11.60, 0.34, fill=LIGHT_ORANGE, line=LIGHT_ORANGE, line_width=0)
    add_text(slide, "AIOps 生产规划/诊断模型要求客户近端 DeepSeek 服务；本地大模型推理建议部署在独立 GPU 节点。", 0.90, 6.62, 11.18, 0.34, size=9.2, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 12, source="来源：Model Serving Provider Catalog（2026-10-12）")


def slide_13(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, INK)
    add_title(slide, "12 · BUSINESS STORY", "一次阻塞告警，如何变成可验证的处理结果", "用一个完整故事理解 KBot AIOps 的客户价值。", dark=True)
    story = [
        ("1", "告警进入", "监控事件验签、映射、去重", BLUE),
        ("2", "自动取证", "阻塞链、活动会话与指标时间线", TEAL),
        ("3", "根因评估", "支持证据、反证与可信等级", BLUE),
        ("4", "形成提案", "风险、影响、动作与验证方案", ORANGE),
        ("5", "人工批准", "逐条审批，Executor 再复核", RED),
        ("6", "验证沉淀", "Before / After 与故障报告", TEAL),
    ]
    for i, (num, title, body, accent) in enumerate(story):
        x = 0.62 + i * 2.09
        add_icon_disc(slide, num, x + 0.55, 2.57, fill=accent, size=0.62, font_size=13)
        if i < len(story) - 1:
            add_line(slide, x + 1.17, 2.88, x + 2.18, 2.88, color="5F7890", width=1.6)
        add_text(slide, title, x, 3.38, 1.72, 0.38, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x, 3.86, 1.72, 0.78, size=9.8, color="B9C8D6", align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
    add_text(slide, "让 AI 参与运维，但让每一步都有证据、有边界、有负责人。", 1.06, 5.08, 11.20, 0.72, size=25, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_round_rect(slide, 3.65, 6.00, 6.05, 0.55, fill=TEAL, line=TEAL, line_width=0)
    add_text(slide, "KBot AIOps · 从“发现异常”走到“证明改善”", 3.90, 6.00, 5.55, 0.55, size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 13, dark=True)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "KBot AIOps：可信、可控的数据库智能运维"
    prs.core_properties.subject = "面向客户和销售人员的 KBot4 AIOps 产品介绍"
    prs.core_properties.author = "KBot"
    prs.core_properties.keywords = "KBot, AIOps, Oracle, MySQL, PostgreSQL, 智能运维"
    prs.core_properties.comments = "基于 2026-10-12 当前 KBot4 代码、契约与部署文档生成。"

    slide_01(prs)
    slide_02(prs)
    slide_03(prs)
    slide_04(prs)
    slide_05(prs)
    slide_06(prs)
    slide_07(prs)
    slide_08(prs)
    slide_09(prs)
    slide_10(prs)
    slide_11(prs)
    slide_12(prs)
    slide_13(prs)

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
